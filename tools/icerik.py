#!/usr/bin/env python3
"""Makarya içerik aracı.

  python3 tools/icerik.py json2xlsx   content/*.json  -> Makarya_Icerik.xlsx   (ilk kurulum)
  python3 tools/icerik.py xlsx2js     Makarya_Icerik.xlsx -> content.js      (her düzenlemeden sonra)
  python3 tools/icerik.py check       Makarya_Icerik.xlsx                    (sadece doğrula)

Excel ana kaynaktır. content.js elle düzenlenmez."""
import glob, json, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "Makarya_Icerik.xlsx")
JS = os.path.join(ROOT, "content.js")

STATS = ["Akıl", "Çene", "Kurnazlık", "Cesaret", "Pişkinlik", "Vicdan", "Dayanıklılık", "Sosyal Radar"]
DIFF_TR = {"veryEasy": "Çok Kolay", "easy": "Kolay", "medium": "Orta", "hard": "Zor", "veryHard": "Çok Zor"}
TR_DIFF = {v: k for k, v in DIFF_TR.items()}
FAM_TR = {"home": "ev", "money": "para", "social": "sosyal", "education": "okul-iş"}
TR_FAM = {v: k for k, v in FAM_TR.items()}
MINI_TYPES = ["timing", "hold", "collect", "race", "memory", "cups", "doors", "rhythm", "lanes", "simon", "poker", "trace",
              "bargain", "swipe", "breath", "balance"]
PLACES = ["ev", "mutfak", "cocuk-odasi", "okul", "sinif", "okul-bahcesi", "sokak", "park", "bakkal", "market", "carsi",
          "ofis", "dukkan", "devlet-dairesi", "hastane", "dugun-salonu", "kafe", "toplu-tasima", "trafik", "universite",
          "kisla", "banka", "apartman", "sahil", "huzurevi", "mezarlik", "yazlik", "stadyum", "sahne", "mahkeme"]
FAMILIES = ["Yalçın", "Erdem", "Keskin", "Tan", "Varlı", "Şen"]

EV_COLS = ["Olay ID", "Bölüm", "Yaş", "Simge", "Başlık", "Olay metni", "Mekân", "Aile türü", "Sabit", "Ağırlık",
           "Gereken hafızalar", "Olmaması gereken hafızalar", "Gereken özellik", "Cinsiyet", "Sadece aileler", "Asgari stat"]
CH_COLS = ["Olay ID", "Sıra", "Seçim metni", "Stat", "Zorluk", "Kesin etki", "Kesin sonuç metni",
           "Başarı", "Kritik başarı", "Başarısız", "Kritik hata", "Yarım (mini oyun)",
           "Mini oyun", "Mini oyun başlığı", "Mini oyun ipucu", "Mini oyun ayarı",
           "Gereken özellik", "Gereken hafıza", "İç ses açar", "Kazandırdığı özellik",
           "Seçince hafıza", "Başarıda hafıza", "Kritik hatada hafıza", "Silinen hafıza", "Ölüm riski %", "Ölüm sebebi"]
VO_COLS = ["Olay ID", "Stat", "Zorluk", "İç ses metni", "Başarısızlık metni", "Seçenek açar"]
RE_COLS = ["Olay ID", "Hafıza", "Geçmişten metni"]
TR_COLS = ["Özellik", "Açıklama", "Bonus", "Otomatik stat"]
FL_COLS = ["Hafıza", "Final kartı etiketi"]
DC_COLS = ["Asgari yaş", "Azami yaş", "Mezar taşı metni", "Gereken hafıza"]
BO_COLS = ["Bölüm", "Ad", "Seçilecek olay", "Beklenen stat", "Yaş aralığı", "Son görsel (olay ID)", "Son yaş", "Son metin", "Sonraki düğme"]


# ── yardımcılar ──
def fmt_stats(d):
    return ", ".join(f"{k} {'+' if v > 0 else ''}{v}" for k, v in (d or {}).items())

def parse_stats(s, where):
    out = {}
    if not s: return out
    for part in re.split(r"[;,]", str(s)):
        part = part.strip()
        if not part: continue
        m = re.match(r"^(.+?)\s*([+-]?\s*\d+)$", part)
        if not m or m.group(1).strip() not in STATS: raise ValueError(f"{where}: stat etkisi okunamadı: '{part}' (ör. 'Vicdan +2, Çene +1')")
        out[m.group(1).strip()] = int(m.group(2).replace(" ", ""))
    return out

def split_list(s):
    return [x.strip() for x in re.split(r"[,;\n]", str(s or "")) if x.strip()]

def yes(v):
    return str(v or "").strip().lower() in ["e", "evet", "x", "1", "true", "✓"]

def clean(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): return int(v)
    return v.strip() if isinstance(v, str) else v


# ── JSON kaynaklarını topla ──
def bundle_from_json(paths):
    b = {"chapters": [], "events": [], "traits": [], "flags": [], "deathCauses": []}
    seen_t, seen_f = set(), set()
    for p in paths:
        d = json.load(open(p, encoding="utf-8"))
        b["chapters"] += d.get("chapters", [])
        b["events"] += d.get("events", [])
        for t in d.get("traits", []):
            if t["name"] not in seen_t: seen_t.add(t["name"]); b["traits"].append(t)
        for f in d.get("flags", []):
            if f["flag"] not in seen_f: seen_f.add(f["flag"]); b["flags"].append(f)
        b["deathCauses"] += d.get("deathCauses", [])
    return b


# ── Excel yaz ──
def write_xlsx(b, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    HEAD = PatternFill("solid", fgColor="1D1915"); HF = Font(bold=True, color="FFFFFF")
    ALT = PatternFill("solid", fgColor="F8F1E3")
    thin = Side(style="thin", color="D8C9AD")

    def sheet(name, cols, rows, widths, wrap_cols=(), first=False):
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        ws.append(cols)
        for r in rows: ws.append([("" if v is None else v) for v in r])
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i); cell.fill = HEAD; cell.font = HF
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cols[cell.column - 1] in wrap_cols)
                cell.border = Border(bottom=thin)
        ws.freeze_panes = "B2" if len(cols) > 3 else "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    def dv(ws, col_name, cols, options, nrows):
        col = get_column_letter(cols.index(col_name) + 1)
        v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True, showErrorMessage=True,
                           errorTitle="Geçersiz değer", error="Listeden bir değer seçin.")
        ws.add_data_validation(v); v.add(f"{col}2:{col}{nrows + 400}")

    # Nasıl kullanılır
    ws = wb.active; ws.title = "Nasıl Kullanılır"
    guide = [
        ["MAKARYA · İÇERİK TABLOSU"],
        [""],
        ["Bu dosya oyunun bütün içeriğinin ana kaynağıdır. Burada yapılan değişiklikler oyuna şöyle geçer:"],
        ["1) Dosyayı düzenle ve kaydet (adını değiştirme: Makarya_Icerik.xlsx)."],
        ["2) GitHub'da depoya yükle (Add file → Upload files) ya da Claude'a gönder."],
        ["3) Dosya oyunun formatına dönüştürülür ve birkaç dakika içinde oyunda görünür."],
        [""],
        ["SAYFALAR"],
        ["Olaylar: Her satır bir olay. Olay ID benzersiz olmalı (küçük harf ve tire). Bölüm 1–10. Yaş örn. '12 yaş 6 ay'."],
        ["Seçimler: Her satır bir seçim; Olay ID ile olaya bağlanır, Sıra ile dizilir. Bir olayda 2–5 seçim olur."],
        ["   • Kesin sonuçlu seçim: 'Kesin etki' (ör. Vicdan +2, Çene +1) ve 'Kesin sonuç metni' doldurulur."],
        ["   • Zarlı seçim: Stat + Zorluk + dört sonuç metni (Başarı, Kritik başarı, Başarısız, Kritik hata)."],
        ["   • Mini oyunlu seçim: zarlı seçim + Mini oyun türü + başlık + ipucu + 'Yarım (mini oyun)' metni."],
        ["   • Ölüm riski yalnızca Bölüm 6 ve sonrası (18+). Yüzde olarak yaz (ör. 3). Oyuncu bu yüzdeyi görür."],
        ["İç Sesler: Olayın başında araya giren stat yorumları. 'Seçenek açar' = E ise, 'İç ses açar' sütununda aynı stat yazan seçim ancak bu ses başarılı olursa görünür."],
        ["Geçmişten: Oyuncunun hafızasında o kayıt varsa olayda gösterilen hatırlatma cümlesi."],
        ["Özellikler: Kazanılabilir özellikler; Bonus ilgili stat kontrollerinde zar hedefini düşürür (ör. Çene +1, Kurnazlık +1)."],
        ["Hafıza: Seçimlerin yazdığı kayıtlar ve final kartındaki açıklamaları."],
        ["Ölüm Sebepleri: 18 yaş sonrası yaşla artan arka plan riskinden ölünce mezar taşına yazılır."],
        ["Bölümler: Her bölümde bir hayatta kaç olay görüleceği (Seçilecek olay), zorluk ayarı (Beklenen stat) ve bölüm sonu ekranı."],
        ["Özet: Bölüm başına olay, mini oyun ve risk sayıları (otomatik hesaplanır)."],
        [""],
        ["KURALLAR (yazım yönergesinin özeti)"],
        ["• Makarya, Türkiye'ye benzeyen kurgusal bir ülkedir ve günümüzde geçer."],
        ["• Marka, uygulama, kulüp, ünlü, gerçek kişi, gerçek şehir/semt/anıt, gerçek sınav ve kurum adı YOK."],
        ["• Din ve siyaset espri konusu değildir; siyasetçi yalnızca isimsiz ve genel olarak geçebilir."],
        ["• Ölüm yalnızca 18 yaşından sonra; şiddet sahne dışında; hastalık, yoksulluk, engellilik alay konusu değildir."],
        ["• Sonuç metinleri 1–2 cümle; espri son kelimede patlasın. Oyuncuya 'sen' diye hitap edilir."],
        ["• Açılır listeli sütunlarda yalnızca listedeki değerleri kullan."],
        [""],
        ["DEĞERLER"],
        ["Stat: " + ", ".join(STATS)],
        ["Zorluk: " + ", ".join(DIFF_TR.values())],
        ["Aile türü: ev, para, sosyal, okul-iş"],
        ["Mini oyun: " + ", ".join(MINI_TYPES)],
        ["Mekân: " + ", ".join(PLACES)],
        ["Aileler: " + ", ".join(FAMILIES)],
    ]
    for r in guide: ws.append(r)
    ws.column_dimensions["A"].width = 150
    ws["A1"].font = Font(bold=True, size=16)
    for r in [8, 24, 31]: ws.cell(row=r, column=1).font = Font(bold=True)

    # Olaylar
    rows_ev, rows_ch, rows_vo, rows_re = [], [], [], []
    for e in b["events"]:
        w = e.get("when") or {}
        rows_ev.append([e["id"], e["ch"], e["age"], e.get("icon", ""), e["title"], e["text"], e.get("place", ""),
                        FAM_TR.get(e.get("fam", ""), ""), "E" if e.get("fixed") else "", e.get("weight", ""),
                        ", ".join(w.get("flags", [])), ", ".join(w.get("notFlags", [])), w.get("trait", ""),
                        {"erkek": "erkek", "kiz": "kız"}.get(w.get("gender", ""), ""), ", ".join(w.get("fam", [])), fmt_stats(w.get("minStat"))])
        for i, c in enumerate(e["choices"], 1):
            r = c.get("r") if isinstance(c.get("r"), dict) else {}
            m = c.get("mini") or {}
            params = {k: v for k, v in m.items() if k not in ["type", "title", "hint"]}
            rk = c.get("risk") or {}
            rows_ch.append([e["id"], i, c["t"], c.get("stat", ""), DIFF_TR.get(c.get("diff", ""), ""), fmt_stats(c.get("direct")),
                            c["r"] if isinstance(c.get("r"), str) else "", r.get("win", ""), r.get("crit", ""), r.get("fail", ""), r.get("bad", ""), r.get("mid", ""),
                            m.get("type", ""), m.get("title", ""), m.get("hint", ""), json.dumps(params, ensure_ascii=False) if params else "",
                            c.get("req", ""), c.get("reqFlag", ""), c.get("reqVoice", ""), c.get("trait", ""),
                            c.get("flag", ""), c.get("flagWin", ""), c.get("flagBad", ""), c.get("unflag", ""),
                            (round(rk["p"] * 100, 2) if rk else ""), rk.get("cause", "")])
        for v in e.get("voices", []):
            rows_vo.append([e["id"], v["stat"], DIFF_TR[v["diff"]], v["text"], v.get("fail", ""), "E" if v.get("opens") else ""])
        for rc in e.get("recall", []):
            rows_re.append([e["id"], rc["flag"], rc["text"]])

    W = {"Olay ID": 26, "Bölüm": 8, "Yaş": 12, "Simge": 7, "Başlık": 26, "Olay metni": 60, "Mekân": 14, "Aile türü": 10, "Sabit": 7,
         "Ağırlık": 8, "Gereken hafızalar": 20, "Olmaması gereken hafızalar": 20, "Gereken özellik": 18, "Cinsiyet": 9, "Sadece aileler": 14, "Asgari stat": 14}
    ws = sheet("Olaylar", EV_COLS, rows_ev, W, wrap_cols=("Olay metni", "Başlık"))
    n = len(rows_ev)
    dv(ws, "Bölüm", EV_COLS, [str(i) for i in range(1, 11)], n); dv(ws, "Mekân", EV_COLS, PLACES, n)
    dv(ws, "Aile türü", EV_COLS, list(TR_FAM.keys()), n); dv(ws, "Sabit", EV_COLS, ["E"], n); dv(ws, "Cinsiyet", EV_COLS, ["erkek", "kız"], n)

    W = {"Olay ID": 26, "Sıra": 6, "Seçim metni": 34, "Stat": 13, "Zorluk": 10, "Kesin etki": 20, "Kesin sonuç metni": 44,
         "Başarı": 44, "Kritik başarı": 44, "Başarısız": 44, "Kritik hata": 44, "Yarım (mini oyun)": 36, "Mini oyun": 11,
         "Mini oyun başlığı": 20, "Mini oyun ipucu": 34, "Mini oyun ayarı": 30, "Gereken özellik": 18, "Gereken hafıza": 18,
         "İç ses açar": 13, "Kazandırdığı özellik": 18, "Seçince hafıza": 16, "Başarıda hafıza": 16, "Kritik hatada hafıza": 16,
         "Silinen hafıza": 14, "Ölüm riski %": 9, "Ölüm sebebi": 44}
    ws = sheet("Seçimler", CH_COLS, rows_ch, W, wrap_cols=("Seçim metni", "Kesin sonuç metni", "Başarı", "Kritik başarı", "Başarısız", "Kritik hata", "Yarım (mini oyun)", "Mini oyun ipucu", "Ölüm sebebi", "Mini oyun ayarı"))
    n = len(rows_ch)
    dv(ws, "Stat", CH_COLS, STATS, n); dv(ws, "Zorluk", CH_COLS, list(TR_DIFF.keys()), n); dv(ws, "Mini oyun", CH_COLS, MINI_TYPES, n)
    dv(ws, "İç ses açar", CH_COLS, STATS, n)

    ws = sheet("İç Sesler", VO_COLS, rows_vo, {"Olay ID": 26, "Stat": 13, "Zorluk": 10, "İç ses metni": 70, "Başarısızlık metni": 40, "Seçenek açar": 9},
               wrap_cols=("İç ses metni", "Başarısızlık metni"))
    dv(ws, "Stat", VO_COLS, STATS, len(rows_vo)); dv(ws, "Zorluk", VO_COLS, ["Kolay", "Orta", "Zor", "Çok Zor"], len(rows_vo)); dv(ws, "Seçenek açar", VO_COLS, ["E"], len(rows_vo))
    sheet("Geçmişten", RE_COLS, rows_re, {"Olay ID": 26, "Hafıza": 20, "Geçmişten metni": 80}, wrap_cols=("Geçmişten metni",))

    ws = sheet("Özellikler", TR_COLS, [[t["name"], t["desc"], fmt_stats(t["bonus"]), t.get("auto", "")] for t in b["traits"]],
               {"Özellik": 26, "Açıklama": 60, "Bonus": 26, "Otomatik stat": 14}, wrap_cols=("Açıklama",))
    dv(ws, "Otomatik stat", TR_COLS, STATS, len(b["traits"]))
    sheet("Hafıza", FL_COLS, [[f["flag"], f["label"]] for f in b["flags"]], {"Hafıza": 26, "Final kartı etiketi": 80}, wrap_cols=("Final kartı etiketi",))
    sheet("Ölüm Sebepleri", DC_COLS, [[d["min"], d["max"], d["text"], d.get("flag", "")] for d in b["deathCauses"]],
          {"Asgari yaş": 10, "Azami yaş": 10, "Mezar taşı metni": 90, "Gereken hafıza": 18}, wrap_cols=("Mezar taşı metni",))
    sheet("Bölümler", BO_COLS, [[c["ch"], c["name"], c["pick"], c["expected"], c["range"], c["endImg"], c["endAge"], c["endText"], c["next"]] for c in b["chapters"]],
          {"Bölüm": 8, "Ad": 18, "Seçilecek olay": 10, "Beklenen stat": 10, "Yaş aralığı": 18, "Son görsel (olay ID)": 22, "Son yaş": 10, "Son metin": 60, "Sonraki düğme": 30},
          wrap_cols=("Son metin",))

    # Özet (formüllerle)
    ws = wb.create_sheet("Özet")
    ws.append(["Bölüm", "Olay sayısı", "Sabit olay", "Bir hayatta görülen", "Seçim sayısı", "Mini oyunlu seçim", "Ölüm riskli seçim"])
    for i in range(1, 11):
        r = i + 1
        ws.append([i, f"=COUNTIF(Olaylar!B:B,{i})", f'=COUNTIFS(Olaylar!B:B,{i},Olaylar!I:I,"E")', f"=VLOOKUP({i},Bölümler!A:C,3,FALSE)",
                   f'=SUMPRODUCT((COUNTIFS(Seçimler!A:A,Olaylar!A$2:A${len(rows_ev)+1})>0)*(Olaylar!B$2:B${len(rows_ev)+1}={i})*COUNTIFS(Seçimler!A:A,Olaylar!A$2:A${len(rows_ev)+1}))',
                   "", ""])
    ws.append(["Toplam", "=SUM(B2:B11)", "=SUM(C2:C11)", "=SUM(D2:D11)", "=SUM(E2:E11)", f'=COUNTIF(Seçimler!M:M,"?*")-1', f'=COUNTIF(Seçimler!Y:Y,">0")'])
    for i, wdt in enumerate([10, 12, 12, 18, 12, 18, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
        ws.cell(row=1, column=i).fill = HEAD; ws.cell(row=1, column=i).font = HF
    ws.cell(row=12, column=1).font = Font(bold=True)

    # sayfa sırası: Nasıl Kullanılır, Özet, Olaylar, Seçimler, ...
    order = ["Nasıl Kullanılır", "Özet", "Olaylar", "Seçimler", "İç Sesler", "Geçmişten", "Özellikler", "Hafıza", "Ölüm Sebepleri", "Bölümler"]
    wb._sheets = [wb[n] for n in order]
    wb.save(path)


# ── Excel oku ──
def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    errs = []

    def rows(name, cols):
        ws = wb[name]
        head = [clean(c.value) for c in ws[1]]
        idx = {}
        for c in cols:
            if c not in head: errs.append(f"'{name}' sayfasında '{c}' sütunu yok"); continue
            idx[c] = head.index(c)
        out = []
        for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not any(v not in (None, "") for v in r): continue
            out.append((i, {c: clean(r[j]) if j < len(r) else "" for c, j in idx.items()}))
        return out

    b = {"chapters": [], "events": [], "traits": [], "flags": [], "deathCauses": []}
    evmap = {}
    for i, r in rows("Olaylar", EV_COLS):
        where = f"Olaylar satır {i}"
        try:
            e = {"id": str(r["Olay ID"]), "ch": int(r["Bölüm"]), "age": str(r["Yaş"]), "icon": str(r["Simge"] or "✦"),
                 "title": str(r["Başlık"]), "text": str(r["Olay metni"])}
        except Exception as ex:
            errs.append(f"{where}: {ex}"); continue
        if r["Mekân"]: e["place"] = r["Mekân"]
        if r["Aile türü"]:
            if r["Aile türü"] not in TR_FAM: errs.append(f"{where}: aile türü '{r['Aile türü']}' geçersiz")
            else: e["fam"] = TR_FAM[r["Aile türü"]]
        if yes(r["Sabit"]): e["fixed"] = True
        if r["Ağırlık"] not in ("", None): e["weight"] = float(r["Ağırlık"])
        w = {}
        if r["Gereken hafızalar"]: w["flags"] = split_list(r["Gereken hafızalar"])
        if r["Olmaması gereken hafızalar"]: w["notFlags"] = split_list(r["Olmaması gereken hafızalar"])
        if r["Gereken özellik"]: w["trait"] = r["Gereken özellik"]
        if r["Cinsiyet"]: w["gender"] = {"erkek": "erkek", "kız": "kiz", "kiz": "kiz"}.get(r["Cinsiyet"], r["Cinsiyet"])
        if r["Sadece aileler"]: w["fam"] = split_list(r["Sadece aileler"])
        if r["Asgari stat"]:
            try: w["minStat"] = parse_stats(r["Asgari stat"], where)
            except ValueError as ex: errs.append(str(ex))
        if w: e["when"] = w
        e["choices"], e["voices"], e["recall"] = [], [], []
        if e["id"] in evmap: errs.append(f"{where}: Olay ID '{e['id']}' tekrar ediyor")
        evmap[e["id"]] = e; b["events"].append(e)

    chs = []
    for i, r in rows("Seçimler", CH_COLS):
        where = f"Seçimler satır {i}"
        e = evmap.get(str(r["Olay ID"]))
        if not e: errs.append(f"{where}: Olay ID '{r['Olay ID']}' Olaylar sayfasında yok"); continue
        c = {"t": str(r["Seçim metni"])}
        try:
            if r["Kesin etki"]:
                c["direct"] = parse_stats(r["Kesin etki"], where); c["r"] = str(r["Kesin sonuç metni"])
            else:
                c["stat"] = r["Stat"]; c["diff"] = TR_DIFF.get(r["Zorluk"], r["Zorluk"])
                c["r"] = {k: str(r[col]) for k, col in [("win", "Başarı"), ("crit", "Kritik başarı"), ("fail", "Başarısız"), ("bad", "Kritik hata"), ("mid", "Yarım (mini oyun)")] if r[col]}
        except ValueError as ex:
            errs.append(str(ex)); continue
        if r["Mini oyun"]:
            m = {"type": r["Mini oyun"], "title": str(r["Mini oyun başlığı"]), "hint": str(r["Mini oyun ipucu"])}
            if r["Mini oyun ayarı"]:
                try: m.update(json.loads(r["Mini oyun ayarı"]))
                except Exception: errs.append(f"{where}: 'Mini oyun ayarı' okunamadı (JSON biçiminde olmalı)")
            c["mini"] = m
        for key, col in [("req", "Gereken özellik"), ("reqFlag", "Gereken hafıza"), ("reqVoice", "İç ses açar"), ("trait", "Kazandırdığı özellik"),
                         ("flag", "Seçince hafıza"), ("flagWin", "Başarıda hafıza"), ("flagBad", "Kritik hatada hafıza"), ("unflag", "Silinen hafıza")]:
            if r[col]: c[key] = str(r[col])
        if r["Ölüm riski %"] not in ("", None):
            c["risk"] = {"p": round(float(r["Ölüm riski %"]) / 100, 4), "cause": str(r["Ölüm sebebi"])}
        chs.append((e, int(r["Sıra"] or 99), i, c))
    for e, _, _, c in sorted(chs, key=lambda x: (x[1], x[2])): e["choices"].append(c)

    for i, r in rows("İç Sesler", VO_COLS):
        e = evmap.get(str(r["Olay ID"]))
        if not e: errs.append(f"İç Sesler satır {i}: Olay ID yok"); continue
        v = {"stat": r["Stat"], "diff": TR_DIFF.get(r["Zorluk"], r["Zorluk"]), "text": str(r["İç ses metni"])}
        if r["Başarısızlık metni"]: v["fail"] = str(r["Başarısızlık metni"])
        if yes(r["Seçenek açar"]): v["opens"] = True
        e["voices"].append(v)
    for i, r in rows("Geçmişten", RE_COLS):
        e = evmap.get(str(r["Olay ID"]))
        if not e: errs.append(f"Geçmişten satır {i}: Olay ID yok"); continue
        e["recall"].append({"flag": str(r["Hafıza"]), "text": str(r["Geçmişten metni"])})
    for e in b["events"]:
        if not e["voices"]: del e["voices"]
        if not e["recall"]: del e["recall"]
    for i, r in rows("Özellikler", TR_COLS):
        try:
            t = {"name": str(r["Özellik"]), "desc": str(r["Açıklama"]), "bonus": parse_stats(r["Bonus"], f"Özellikler satır {i}")}
        except ValueError as ex: errs.append(str(ex)); continue
        if r["Otomatik stat"]: t["auto"] = r["Otomatik stat"]
        b["traits"].append(t)
    for i, r in rows("Hafıza", FL_COLS): b["flags"].append({"flag": str(r["Hafıza"]), "label": str(r["Final kartı etiketi"])})
    for i, r in rows("Ölüm Sebepleri", DC_COLS):
        d = {"min": int(r["Asgari yaş"]), "max": int(r["Azami yaş"]), "text": str(r["Mezar taşı metni"])}
        if r["Gereken hafıza"]: d["flag"] = str(r["Gereken hafıza"])
        b["deathCauses"].append(d)
    for i, r in rows("Bölümler", BO_COLS):
        b["chapters"].append({"ch": int(r["Bölüm"]), "name": str(r["Ad"]), "pick": int(r["Seçilecek olay"]), "expected": int(r["Beklenen stat"]),
                              "range": str(r["Yaş aralığı"]), "endImg": str(r["Son görsel (olay ID)"] or ""), "endAge": str(r["Son yaş"] or ""),
                              "endText": str(r["Son metin"] or ""), "next": str(r["Sonraki düğme"] or "")})
    return b, errs


# ── content.js yaz ──
def write_js(b, path):
    chapters = {c["ch"]: {k: v for k, v in c.items() if k != "ch"} for c in b["chapters"]}
    traits = {t["name"]: {k: v for k, v in t.items() if k != "name"} for t in b["traits"]}
    labels = {f["flag"]: f["label"] for f in b["flags"]}
    J = lambda o: json.dumps(o, ensure_ascii=False, separators=(",", ":"))
    with open(path, "w", encoding="utf-8") as f:
        f.write("/* OTOMATİK ÜRETİLDİ — elle düzenleme. Kaynak: Makarya_Icerik.xlsx · tools/icerik.py xlsx2js */\n")
        f.write("var CHAPTERS = " + J(chapters) + ";\n")
        f.write("var TRAITS = " + J(traits) + ";\n")
        f.write("var FLAG_LABELS = " + J(labels) + ";\n")
        f.write("var DEATH_CAUSES = " + J(b["deathCauses"]) + ";\n")
        f.write("var EVENTS = [\n" + ",\n".join(J(e) for e in b["events"]) + "\n];\n")


def validate_bundle(b):
    import subprocess, tempfile
    tmp = tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8")
    json.dump({k: b[k] for k in ["events", "traits", "flags", "deathCauses"]}, tmp, ensure_ascii=False); tmp.close()
    r = subprocess.run([sys.executable, os.path.join(ROOT, "tools", "validate_content.py"), tmp.name], capture_output=True, text=True)
    os.unlink(tmp.name)
    return r.returncode, r.stdout.replace(tmp.name, "Excel")


def main():
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "json2xlsx":
        paths = sys.argv[2:] or sorted(glob.glob(os.path.join(ROOT, "content", "ilk-taslak", "*.json")))
        b = bundle_from_json(paths)
        def months(a):
            y = re.search(r"(\d+)\s*yaş", a); m = re.search(r"(\d+)\s*ay", a)
            return (int(y.group(1)) * 12 if y else 0) + (int(m.group(1)) if m else 0)
        b["events"].sort(key=lambda e: (e["ch"], months(e["age"])))
        b["chapters"].sort(key=lambda c: c["ch"])
        write_xlsx(b, XLSX); print(f"{XLSX} yazıldı: {len(b['events'])} olay")
    elif cmd in ("xlsx2js", "check"):
        path = sys.argv[2] if len(sys.argv) > 2 else XLSX
        b, errs = read_xlsx(path)
        if errs:
            print("\n".join("HATA  " + e for e in errs)); sys.exit(1)
        code, out = validate_bundle(b)
        print(out.strip().split("\n")[-1] if code == 0 else out)
        if code != 0: sys.exit(1)
        if cmd == "xlsx2js":
            write_js(b, JS); print(f"{JS} yazıldı")
    else:
        print(__doc__)

if __name__ == "__main__":
    main()
