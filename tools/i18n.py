#!/usr/bin/env python3
"""Dil aracı. Kullanım:
  python3 tools/i18n.py export   -> Makarya_Ceviri.xlsx (TR, EN sütunları; mevcut çeviriler korunur)
  python3 tools/i18n.py build    -> lang/en.js üretir (dolu EN hücreleri)
  python3 tools/i18n.py durum    -> kaç satır çevrildi
Sayfalar: Icerik (olay/aile/özellik metinleri), Arayuz (T("...") metinleri)."""
import sys, json, re, subprocess, os
from openpyxl import Workbook, load_workbook
os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
X = "Makarya_Ceviri.xlsx"
JS = """global.localStorage={getItem:()=>"tr"};
const fs=require("fs"),vm=require("vm");
vm.runInThisContext(["data.js","content.js","i18n.js"].map(f=>fs.readFileSync(f,"utf8")).join("\\n")+"\\nconsole.log(JSON.stringify(i18nCollect()));");"""
def icerik():
    return json.loads(subprocess.check_output(["node", "-e", JS]))
def ui():
    out = []
    for f in ("game.js", "mini.js", "rules.js"):
        out += re.findall(r'\bT\("((?:[^"\\]|\\.)*)"\)', open(f, encoding="utf8").read())
    return sorted({json.loads('"%s"' % s) for s in out})
def mevcut():
    d = {"Icerik": {}, "Arayuz": {}}
    if os.path.exists(X):
        wb = load_workbook(X)
        for n in d:
            if n in wb.sheetnames:
                for r in wb[n].iter_rows(min_row=2, values_only=True):
                    if r[0] and r[1]: d[n][r[0]] = r[1]
    return d
def export():
    m = mevcut(); wb = Workbook(); wb.remove(wb.active)
    for n, keys in (("Icerik", icerik()), ("Arayuz", ui())):
        ws = wb.create_sheet(n); ws.append(["TR", "EN"])
        for k in keys: ws.append([k, m[n].get(k, "")])
        ws.column_dimensions["A"].width = 80; ws.column_dimensions["B"].width = 80
    wb.save(X); print("yazıldı:", X)
def build():
    m = mevcut(); js = lambda d: json.dumps(d, ensure_ascii=False, separators=(",", ":"))
    open("lang/en.js", "w", encoding="utf8").write("/* tools/i18n.py build ile üretilir — elle düzenleme, Makarya_Ceviri.xlsx'i düzenle */\nEN_DICT = %s;\nEN_UI = %s;\n" % (js(m["Icerik"]), js(m["Arayuz"])))
    print("lang/en.js:", len(m["Icerik"]), "içerik,", len(m["Arayuz"]), "arayüz")
def durum():
    m = mevcut(); print("içerik %d/%d, arayüz %d/%d" % (len(m["Icerik"]), len(icerik()), len(m["Arayuz"]), len(ui())))
{"export": export, "build": build, "durum": durum}.get(sys.argv[1] if len(sys.argv) > 1 else "", lambda: print(__doc__))()
